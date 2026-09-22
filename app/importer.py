"""엑셀(xlsm/xlsx) -> SQLite 임포트.

참고 워크북의 'raw' 시트(기업별 계정 원장)와 '정보입력' 시트(기업 마스터, 계정 매핑)를 읽는다.
"""
import re
import warnings
from pathlib import Path

from . import categories as C
from .db import get_conn, now

warnings.filterwarnings("ignore", module="openpyxl")

YEAR_RE = re.compile(r"(20\d{2})")


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _s(v):
    if v is None:
        return None
    s = str(v).replace("\xa0", " ").strip()
    return s or None


def parse_workbook(path):
    """워크북을 읽어 (companies, account_map, facts) 반환."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    companies = {}
    account_map = {}
    facts = {}

    # ---- 정보입력: 기업 마스터 + 계정 매핑
    if "정보입력" in wb.sheetnames:
        ws = wb["정보입력"]
        rows = ws.iter_rows(values_only=True)
        header = [(_s(h) or "") for h in next(rows)]
        col = {h: i for i, h in enumerate(header)}
        name_i = col.get("기업명", 1)
        sector_i = col.get("대분류", 2)
        cat_i = col.get("중분류", 3)
        desc_i = col.get("기업설명", 4)
        inc_i = col.get("업종합계대상", 11)
        dart_i = col.get("DART")
        map_i = col.get("판관비 항목")
        for r in rows:
            name = _s(r[name_i]) if name_i < len(r) else None
            if name:
                companies[name] = {
                    "name": name,
                    "sector": _s(r[sector_i]) if sector_i < len(r) else None,
                    "category": _s(r[cat_i]) if cat_i < len(r) else None,
                    "description": _s(r[desc_i]) if desc_i < len(r) else None,
                    "include_in_sector": 0 if (_s(r[inc_i]) if inc_i < len(r) else "O") == "X" else 1,
                }
            if dart_i is not None and map_i is not None and dart_i < len(r) and map_i < len(r):
                src, cat = _s(r[dart_i]), C.normalize(r[map_i])
                if src and cat:
                    account_map[src] = cat

    # ---- raw: 계정 원장
    if "raw" in wb.sheetnames:
        ws = wb["raw"]
        rows = ws.iter_rows(values_only=True)
        header = [(_s(h) or "") for h in next(rows)]
        col = {h: i for i, h in enumerate(header)}
        name_i = col.get("기업명", 6)
        cat_i = col.get("계정분류", 7)
        item_i = col.get("최초계정", 11)
        sector_i = col.get("업종", 1)
        sub_i = col.get("카테고리", 2)
        inc_i = col.get("업종합계대상", 0)
        year_cols = [(int(YEAR_RE.search(h).group(1)), i) for h, i in col.items() if h and YEAR_RE.fullmatch(h.replace("년", ""))]
        for r in rows:
            name = _s(r[name_i]) if name_i < len(r) else None
            cat = C.normalize(r[cat_i]) if cat_i < len(r) else None
            if not name or not cat:
                continue
            if name not in companies:
                companies[name] = {
                    "name": name,
                    "sector": _s(r[sector_i]) if sector_i < len(r) else None,
                    "category": _s(r[sub_i]) if sub_i < len(r) else None,
                    "description": None,
                    "include_in_sector": 0 if (_s(r[inc_i]) if inc_i < len(r) else "O") == "X" else 1,
                }
            item = (_s(r[item_i]) if item_i < len(r) else None) or ""
            for year, yi in year_cols:
                amt = _num(r[yi]) if yi < len(r) else None
                if amt is None:
                    continue
                key = (name, cat, item, year)
                facts[key] = facts.get(key, 0.0) + amt
    wb.close()
    return companies, account_map, facts


def write_to_db(companies, account_map, facts, source="excel", replace_company_facts=True):
    """파싱 결과를 DB에 반영. 반환: 기록된 fact 수.

    replace_company_facts=True 면 facts 에 등장하는 기업의 기존 데이터를 (소스와 무관하게) 모두 지우고 새로 기록한다.
    엑셀 원장은 해당 기업의 전체 데이터이므로 오래된 행이 남지 않도록 하기 위함."""
    ts = now()
    with get_conn() as conn:
        for c in companies.values():
            conn.execute(
                """INSERT INTO companies(name, sector, category, description, include_in_sector, updated_at)
                   VALUES(?,?,?,?,?,?)
                   ON CONFLICT(name) DO UPDATE SET
                     sector=COALESCE(excluded.sector, sector),
                     category=COALESCE(excluded.category, category),
                     description=COALESCE(excluded.description, description),
                     include_in_sector=excluded.include_in_sector,
                     updated_at=excluded.updated_at""",
                (c["name"], c["sector"], c["category"], c["description"], c["include_in_sector"], ts),
            )
        for src, cat in account_map.items():
            conn.execute(
                "INSERT INTO account_map(source_name, category) VALUES(?,?) ON CONFLICT(source_name) DO UPDATE SET category=excluded.category",
                (src, cat),
            )
        ids = {r["name"]: r["id"] for r in conn.execute("SELECT id, name FROM companies")}
        if replace_company_facts:
            touched = {ids[n] for (n, _, _, _) in facts if n in ids}
            for cid in touched:
                conn.execute("DELETE FROM facts WHERE company_id=?", (cid,))
        n = 0
        for (name, cat, item, year), amt in facts.items():
            conn.execute(
                """INSERT INTO facts(company_id, category, item, fiscal_year, amount, source, updated_at)
                   VALUES(?,?,?,?,?,?,?)
                   ON CONFLICT(company_id, category, item, fiscal_year) DO UPDATE SET
                     amount=excluded.amount, source=excluded.source, updated_at=excluded.updated_at""",
                (ids[name], cat, item, year, amt, source, ts),
            )
            n += 1
    return n


def import_file(path):
    path = Path(path)
    companies, account_map, facts = parse_workbook(path)
    n = write_to_db(companies, account_map, facts, source="excel")
    return {"companies": len(companies), "account_map": len(account_map), "facts": n}


if __name__ == "__main__":
    import sys
    from .db import init_db

    init_db()
    for p in sys.argv[1:]:
        print(p, import_file(p))
